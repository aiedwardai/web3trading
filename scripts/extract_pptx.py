#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
web3trading / PPTX 课程材料提取

用法：
  python extract_pptx.py <file.pptx> [输出目录]

产出（输出到目录）：
  slides_text.txt       逐页文本（标题/正文/表格/演讲者备注），带页码
  pics/s{NN}_{i}.png    内嵌图片（图表通常在这里，纯图片页必须目视核对）
  pics_manifest.txt     图片 -> 页码映射 + 文件大小（按大小排序可快速定位主图）

依赖：python-pptx（已在 C:/Users/Edward/.workbuddy/binaries/python/envs/default/... 装好）
"""

import os
import pathlib
import sys

from pptx import Presentation


def dump_text(prs, out):
    lines = [f"幻灯片数: {len(prs.slides)}", f"尺寸: {prs.slide_width} x {prs.slide_height}"]
    for i, s in enumerate(prs.slides, 1):
        lines.append("\n" + "=" * 70)
        layout = s.slide_layout.name if s.slide_layout is not None else "?"
        lines.append(f"[SLIDE {i}] layout={layout}")
        for sh in s.shapes:
            try:
                if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                    t = sh.text_frame.text.strip()
                    if t:
                        lines.append(f"  <{sh.shape_type}> {t}")
                if getattr(sh, "has_table", False) and sh.has_table:
                    lines.append("  [TABLE]")
                    for r in sh.table.rows:
                        lines.append("   | " + " | ".join(c.text.strip() for c in r.cells))
            except Exception as e:  # noqa: BLE001
                lines.append(f"  [ERR] {e}")
        if s.has_notes_slide:
            nt = s.notes_slide.notes_text_frame.text.strip()
            if nt:
                lines.append(f"  [NOTES] {nt}")
    out.write_text("\n".join(lines), encoding="utf-8")


def dump_pics(prs, pics_dir):
    pics_dir.mkdir(parents=True, exist_ok=True)
    items = []
    for i, s in enumerate(prs.slides, 1):
        idx = 0

        def walk(shapes):
            nonlocal idx
            for sh in shapes:
                # 13 = PICTURE
                if sh.shape_type == 13:
                    try:
                        fn = pics_dir / f"s{i:02d}_{idx}.{sh.image.ext}"
                        fn.write_bytes(sh.image.blob)
                        items.append((fn.name, i, fn.stat().st_size))
                        idx += 1
                    except Exception as e:  # noqa: BLE001
                        print(f"  [pic err] slide {i}: {e}", file=sys.stderr)
                # 6 = GROUP，需递归才能拿到组内图片
                if sh.shape_type == 6:
                    try:
                        walk(sh.shapes)
                    except Exception:  # noqa: BLE001
                        pass

        walk(s.shapes)
    items.sort(key=lambda x: -x[2])
    man = ["图片清单（按大小降序，主图通常靠前）", ""]
    for name, slide, size in items:
        man.append(f"  {name:<20} slide {slide:>3}   {size/1024:>8.0f} KB")
    (pics_dir.parent / "pics_manifest.txt").write_text("\n".join(man), encoding="utf-8")
    return items


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else "pptx_extract"
    out = os.path.abspath(out_dir)
    os.makedirs(out, exist_ok=True)

    prs = Presentation(src)
    dump_text(prs, __import__("pathlib").Path(out) / "slides_text.txt")
    import pathlib
    items = dump_pics(prs, pathlib.Path(out) / "pics")
    print(f"[OK] 文本 -> {out}/slides_text.txt")
    print(f"[OK] 图片 {len(items)} 张 -> {out}/pics/")
    print(f"[OK] 清单 -> {out}/pics_manifest.txt")
    print("\n提示：纯图片页（文本层为空）必须逐页目视核对，不能只靠文本层。")


if __name__ == "__main__":
    main()
